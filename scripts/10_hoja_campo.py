"""Hoja de campo imprimible y GPX de los diez puntos (protocolo seccion 20.5).

    python scripts/10_hoja_campo.py

Solo usa resultados/campo/puntos.csv (los puntos barajados, sin roles): la hoja no puede
delatar que punto es sitio y cual control. Todos los puntos se dibujan y se describen igual.
Escribe resultados/campo/hoja_campo.html, hoja_campo.pdf (con Chrome sin ventana) y puntos.gpx.
"""
import os, csv, glob, base64, io, pathlib, subprocess
import numpy as np
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.colors import ListedColormap

CAMPO = os.path.join("resultados", "campo")
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
ZOOM = (23.105, 23.165, -82.368, -82.322)       # lat min, lat max, lon min, lon max


def mapa(puntos, lat, lon, agua, caja, titulo, marco=None):
    fig, ax = plt.subplots(figsize=(7.2, 7.2 * (caja[1] - caja[0]) / ((caja[3] - caja[2]) * 0.92)))
    ax.pcolormesh(lon, lat, agua.astype(float), cmap=ListedColormap(["#ece8df", "#bcd7ea"]),
                  shading="auto", rasterized=True)
    for p, la, lo in puntos:
        if caja[0] <= la <= caja[1] and caja[2] <= lo <= caja[3]:
            ax.plot(lo, la, "o", ms=7, mfc="#c0392b", mec="white", mew=1.2)
            # si hay otro punto con codigo mayor muy cerca, esta etiqueta va abajo a la izquierda
            cerca = any(q > p and abs(a - la) < 0.006 and abs(b - lo) < 0.006 for q, a, b in puntos)
            ax.annotate(p, (lo, la), xytext=(-30, -15) if cerca and marco else (5, 5),
                        textcoords="offset points", fontsize=10,
                        weight="bold", bbox=dict(boxstyle="round,pad=0.15", fc="white", ec="none", alpha=0.8))
    if marco:
        ax.add_patch(plt.Rectangle((marco[2], marco[0]), marco[3] - marco[2], marco[1] - marco[0],
                                   fill=False, ls="--", lw=1, ec="#555"))
    km = 1 / (111.32 * np.cos(np.radians(23.13)))   # 1 km en grados de longitud
    x0, y0 = caja[2] + 0.04 * (caja[3] - caja[2]), caja[0] + 0.05 * (caja[1] - caja[0])
    ax.plot([x0, x0 + km], [y0, y0], "k-", lw=3)
    ax.text(x0 + km / 2, y0 + 0.012 * (caja[1] - caja[0]), "1 km", ha="center", fontsize=9)
    ax.set_xlim(caja[2], caja[3]); ax.set_ylim(caja[0], caja[1])
    ax.set_aspect(1 / np.cos(np.radians(23.13)))
    ax.set_xlabel("longitud (°)"); ax.set_ylabel("latitud (°)"); ax.set_title(titulo, fontsize=11)
    ax.tick_params(labelsize=8); ax.ticklabel_format(useOffset=False)
    buf = io.BytesIO(); fig.savefig(buf, format="png", dpi=170, bbox_inches="tight"); plt.close(fig)
    return base64.b64encode(buf.getvalue()).decode()


FICHA = """
<div class="ficha"><h3>{p} &nbsp;·&nbsp; {la:.4f}, {lo:.4f}</h3>
<p>Fecha ________ &nbsp; Hora de llegada ______ &nbsp; Visita nº (1, 2 o 3) ___ &nbsp; Nº de la foto de esta hoja ______</p>
<p><b>Dónde estoy</b> (GPS del móvil): lat ______________ lon ______________ &nbsp; Captura de Google Maps: foto nº ______</p>
<p><b>Acceso:</b> □ desde tierra &nbsp; □ desde el agua &nbsp; □ no accesible → ¿por qué? ____________________</p>
<p><b>¿Llovió en los 2 días anteriores?</b> □ no &nbsp; □ poco &nbsp; □ mucho &nbsp; □ no sé</p>
<p><b>1. ¿Hay algo que vierta agua a la bahía a menos de unos 150 m?</b> □ sí &nbsp; □ no &nbsp; □ no se puede ver<br>
Si sí: □ tubo &nbsp; □ alcantarilla o colector &nbsp; □ canal o zanja &nbsp; □ río o arroyo &nbsp; □ otro ____________<br>
¿Está saliendo agua ahora? □ sí &nbsp; □ no &nbsp; □ no se ve &nbsp;&nbsp; ¿A cuántos metros del punto? _____</p>
<p><b>2. El agua:</b> color □ normal &nbsp; □ marrón o turbia &nbsp; □ verde &nbsp; □ gris o negra &nbsp; □ otro ______<br>
□ mancha o pluma &nbsp; □ espuma &nbsp; □ película o aceite &nbsp; □ basura flotando &nbsp; □ nada de esto<br>
<b>Basura que tapa el agua de frente</b> (hasta ~50 m): □ ninguna &nbsp; □ poca (menos de 1 de cada 10 partes) &nbsp; □ bastante (1 a 5 de 10) &nbsp; □ mucha (más de 5 de 10)<br>
Olor: □ ninguno &nbsp; □ a cloaca &nbsp; □ a químico o combustible &nbsp; □ otro ______</p>
<p><b>3. Fotos</b> de la nº ______ a la nº ______ &nbsp;&nbsp; <b>Notas:</b></p>
<div class="notas"></div></div>"""


FILAS_EXCEL = [  # (campo, opciones de la lista desplegable o None)
    ("Punto", None), ("Latitud, longitud", None), ("Fecha", None), ("Hora de llegada", None),
    ("Visita nº", "1,2,3"), ("Dónde estoy: latitud, longitud (GPS del móvil)", None),
    ("Nº de la foto de la captura de Google Maps", None), ("Nº de la foto de esta hoja", None),
    ("Acceso", "desde tierra,desde el agua,no accesible"), ("Si no accesible, ¿por qué?", None),
    ("¿Llovió en los 2 días anteriores?", "no,poco,mucho,no sé"),
    ("1. ¿Hay algo que vierta agua a la bahía a menos de unos 150 m?", "sí,no,no se puede ver"),
    ("Si sí: tipo de salida", "tubo,alcantarilla o colector,canal o zanja,río o arroyo,otro"),
    ("Si «otro», ¿cuál?", None), ("¿Está saliendo agua ahora?", "sí,no,no se ve"),
    ("¿A cuántos metros del punto?", None),
    ("2. El agua: color", "normal,marrón o turbia,verde,gris o negra,otro"), ("Si «otro», ¿cuál? (color)", None),
    ("Mancha o pluma", "☐,☒"), ("Espuma", "☐,☒"), ("Película o aceite", "☐,☒"), ("Basura flotando", "☐,☒"),
    ("Nada de esto", "☐,☒"),
    ("Basura que tapa el agua de frente (hasta ~50 m)",
     "ninguna,poca (menos de 1 de 10),bastante (1 a 5 de 10),mucha (más de 5 de 10)"),
    ("Olor", "ninguno,a cloaca,a químico o combustible,otro"), ("Si «otro», ¿cuál? (olor)", None),
    ("3. Fotos de la nº", None), ("Fotos hasta la nº", None), ("Notas", None)]


def fichas_excel(puntos, carpeta):
    """Una ficha .xlsx por punto, con la misma estructura que las que usó la persona de campo
    en la visita 1 (Campo | Respuesta, listas desplegables) y las dos filas nuevas del 26/10."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    os.makedirs(carpeta, exist_ok=True)
    for p, la, lo in puntos:
        wb = Workbook(); ws = wb.active; ws.title = f"Ficha {p}"
        ws["A1"] = f"Ficha de campo · {p} · {la:.4f}, {lo:.4f}"; ws["A1"].font = Font(bold=True, size=13)
        ws.merge_cells("A1:C1")
        for c, t in (("A4", "Campo"), ("B4", "Respuesta")):
            ws[c] = t; ws[c].font = Font(bold=True, color="FFFFFF"); ws[c].fill = PatternFill("solid", fgColor="1F4E78")
        for i, (campo, opciones) in enumerate(FILAS_EXCEL, start=5):
            ws.cell(i, 1, campo).font = Font(bold=True)
            ws.cell(i, 1).fill = PatternFill("solid", fgColor="D9E1F2")
            if opciones:
                dv = DataValidation(type="list", formula1=f'"{opciones}"', allow_blank=True)
                ws.add_data_validation(dv); dv.add(ws.cell(i, 2))
                if opciones == "☐,☒":
                    ws.cell(i, 2, "☐")
        ws["B5"], ws["B6"] = p, f"{la:.4f}, {lo:.4f}"
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 52, 40
        wb.save(os.path.join(carpeta, f"Ficha_{p}.xlsx"))


def main():
    puntos = [(r["punto"], float(r["lat"]), float(r["lon"]))
              for r in csv.DictReader(open(os.path.join(CAMPO, "puntos.csv")))]
    agua = np.load("resultados/climatologia.npz")["agua"]
    with np.load(sorted(glob.glob("datos/productos/*.npz"))[0]) as z:
        lat, lon = z["lat"], z["lon"]
    la = [p[1] for p in puntos]; lo = [p[2] for p in puntos]
    caja = (min(la) - 0.012, max(la) + 0.012, min(lo) - 0.012, max(lo) + 0.012)
    m1 = mapa(puntos, lat, lon, agua, caja, "Los diez puntos (el recuadro es el mapa siguiente)", ZOOM)
    m2 = mapa(puntos, lat, lon, agua, ZOOM, "Detalle de la bahía")
    tabla = "".join(f"<tr><td><b>{p}</b></td><td>{a:.4f}</td><td>{b:.4f}</td></tr>" for p, a, b in puntos)
    fichas = "".join(FICHA.format(p=p, la=a, lo=b) + ('<div class="salto"></div>' if i % 3 == 2 else "")
                     for i, (p, a, b) in enumerate(puntos))
    html = f"""<!doctype html><html lang="es"><head><meta charset="utf-8"><title>Hoja de campo</title>
<style>
@page {{ size: A4; margin: 14mm; }}
body {{ font-family: Helvetica, Arial, sans-serif; font-size: 10.5pt; color: #111; line-height: 1.35; }}
h1 {{ font-size: 16pt; margin: 0 0 4px; }} h2 {{ font-size: 12.5pt; margin: 14px 0 4px; }}
h3 {{ font-size: 11.5pt; margin: 0 0 6px; }}
li {{ margin-bottom: 3px; }} img {{ width: 100%; }}
table {{ border-collapse: collapse; margin-top: 6px; }} td, th {{ border: 1px solid #999; padding: 2px 12px; }}
.salto {{ break-after: page; }}
.ficha {{ border: 1.5px solid #333; border-radius: 4px; padding: 8px 10px; margin-bottom: 10px; break-inside: avoid; }}
.ficha p {{ margin: 4px 0; }} .notas {{ height: 40px; border-bottom: 1px dotted #888; }}
</style></head><body>
<h1>Visita de campo · Bahía de La Habana</h1>
<p>Es un estudio de la calidad del agua de la bahía con imágenes de satélite. Necesitamos saber
qué hay en la orilla en estos diez puntos: si hay salidas de agua hacia la bahía (tubos,
alcantarillas, canales, ríos) y cómo se ve el agua ese día.</p>
<h2>Cómo hacerlo</h2>
<ul>
<li>Ve a los diez puntos y haz <b>lo mismo en todos, con el mismo cuidado</b>. Que en un punto no haya nada es un dato tan útil como encontrar algo.</li>
<li>Mejor entre las <b>11:00 y las 13:00</b>, que es cuando pasa el satélite.</li>
<li>Activa la <b>ubicación precisa</b> del móvil y la ubicación de la cámara antes de empezar.</li>
<li>En cada punto, abre Google Maps, mantén pulsado el punto azul (donde estás) para ver tus coordenadas, cópialas en «Dónde estoy» y haz una <b>captura de pantalla</b>.</li>
<li><b>Basura que tapa el agua:</b> mira el agua de frente, desde la orilla hasta unos 50 m (media cuadra). Imagina esa superficie partida en 10 trozos iguales y cuenta cuántos tapa la basura flotante. Ninguno: «ninguna»; menos de uno: «poca»; de 1 a 5: «bastante»; más de 5: «mucha». Es a ojo: no hace falta afinar más.</li>
<li>Al llegar a cada punto, haz primero una foto de su ficha (para saber a qué punto son las fotos que siguen). Después: la orilla hacia un lado, la orilla hacia el otro, el agua de frente, y de cerca cualquier tubo, salida, mancha o espuma que veas.</li>
<li>Recorre la orilla <b>unos 150 m a cada lado</b> del punto (una cuadra y media, más o menos). Algunos puntos están en el agua, cerca de la orilla: mira la orilla más cercana.</li>
<li>Rellena la ficha en el momento, no después.</li>
</ul>
<h2>Seguridad</h2>
<ul>
<li>No entres en zonas cerradas, portuarias o militares sin permiso. Si te dicen que no se puede pasar o fotografiar, no insistas: marca «no accesible» y sigue.</li>
<li>No toques el agua ni cojas muestras. Mejor no ir solo a sitios aislados.</li>
</ul>
<h2>Cómo enviarlo</h2>
<ul>
<li>Las fotos, como <b>archivo o documento</b> (o por Google Drive), <b>no</b> como foto normal de WhatsApp, que borra la ubicación y la hora.</li>
<li>Una foto de cada ficha rellena.</li>
</ul>
<p>Esta es la primera de tres visitas: repetiremos en la temporada seca y justo después de un aguacero fuerte.</p>
<div class="salto"></div>
<img src="data:image/png;base64,{m1}">
<table><tr><th>Punto</th><th>Latitud</th><th>Longitud</th></tr>{tabla}</table>
<p style="font-size:9pt">Las coordenadas se pueden pegar en Google Maps (por ejemplo «23.1470, -82.3908»). El archivo <i>puntos.gpx</i> se abre en Maps.me u OsmAnd.</p>
<div class="salto"></div>
<img src="data:image/png;base64,{m2}">
<div class="salto"></div>
{fichas}
</body></html>"""
    ruta = os.path.join(CAMPO, "hoja_campo.html")
    open(ruta, "w").write(html)
    gpx = "".join(f'<wpt lat="{a:.4f}" lon="{b:.4f}"><name>{p}</name></wpt>\n' for p, a, b in puntos)
    open(os.path.join(CAMPO, "puntos.gpx"), "w").write(
        '<?xml version="1.0" encoding="UTF-8"?>\n<gpx version="1.1" creator="Vertidos-recurrentes" '
        'xmlns="http://www.topografix.com/GPX/1/1">\n' + gpx + "</gpx>\n")
    fichas_excel(puntos, os.path.join(CAMPO, "fichas_excel"))
    pdf = os.path.abspath(os.path.join(CAMPO, "hoja_campo.pdf"))
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={pdf}", pathlib.Path(os.path.abspath(ruta)).as_uri()],
                   check=True, capture_output=True, timeout=120)
    print("escrito", pdf)


if __name__ == "__main__":
    main()
